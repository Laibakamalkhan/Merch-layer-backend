from __future__ import annotations

import csv
import datetime
import io
import json
from typing import Any, Iterator

from commerce_connector.platform.datasets.validator import SecurityValidationError, sanitize_csv_cell

MAX_SHEETS_LIMIT = 10
MAX_CELLS_LIMIT = 500_000
MAX_JSON_DEPTH = 5


def _check_json_depth(obj: Any, current_depth: int = 1) -> None:
    if current_depth > MAX_JSON_DEPTH:
        raise SecurityValidationError(f"JSON nesting depth exceeds maximum allowed limit of {MAX_JSON_DEPTH}")
    if isinstance(obj, dict):
        for v in obj.values():
            if isinstance(v, (dict, list)):
                _check_json_depth(v, current_depth + 1)
    elif isinstance(obj, list):
        for v in obj:
            if isinstance(v, (dict, list)):
                _check_json_depth(v, current_depth + 1)


def parse_csv_stream(data: bytes, *, max_rows: int | None = None) -> Iterator[dict[str, Any]]:
    """Parse CSV with encoding detection, delimiter sniffing, and cell sanitization."""
    # Detect encoding
    text = ""
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if not text:
        raise SecurityValidationError("Could not decode CSV using UTF-8 or Latin-1 encoding")

    # Detect delimiter
    sample = "\n".join(text.splitlines()[:20])
    delimiter = ","
    try:
        sniffer = csv.Sniffer()
        dialect = sniffer.sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except Exception:
        delimiter = ","

    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    if not reader.fieldnames:
        return

    count = 0
    for row in reader:
        sanitized_row: dict[str, Any] = {}
        for k, v in row.items():
            if k is None:
                continue
            clean_key = str(k).strip()
            if clean_key:
                sanitized_row[clean_key] = sanitize_csv_cell(v)
        yield sanitized_row
        count += 1
        if max_rows is not None and count >= max_rows:
            break


def parse_xlsx_stream(data: bytes, *, max_rows: int | None = None) -> Iterator[dict[str, Any]]:
    """Parse XLSX spreadsheet with streaming read_only mode and cell limits."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except ImportError:
        raise SecurityValidationError("openpyxl is required to parse XLSX spreadsheets")
    except Exception as exc:
        raise SecurityValidationError(f"Failed to parse XLSX spreadsheet: {exc}") from exc

    sheet_names = wb.sheetnames
    if len(sheet_names) > MAX_SHEETS_LIMIT:
        wb.close()
        raise SecurityValidationError(f"XLSX contains {len(sheet_names)} sheets (limit is {MAX_SHEETS_LIMIT})")

    ws = wb.active or wb[sheet_names[0]]
    headers: list[str] = []
    total_cells = 0
    row_count = 0

    try:
        for row in ws.iter_rows(values_only=True):
            if not row or all(c is None for c in row):
                continue
            total_cells += len(row)
            if total_cells > MAX_CELLS_LIMIT:
                raise SecurityValidationError(f"XLSX cell count exceeded maximum limit of {MAX_CELLS_LIMIT}")

            if not headers:
                headers = [str(c).strip() if c is not None else f"col_{idx}" for idx, c in enumerate(row)]
                continue

            row_dict: dict[str, Any] = {}
            for col_idx, val in enumerate(row):
                if col_idx < len(headers):
                    key = headers[col_idx]
                    if isinstance(val, (datetime.datetime, datetime.date)):
                        val = val.isoformat()
                    row_dict[key] = sanitize_csv_cell(val)
            yield row_dict
            row_count += 1
            if max_rows is not None and row_count >= max_rows:
                break
    finally:
        wb.close()


def parse_json_stream(data: bytes, *, max_rows: int | None = None) -> Iterator[dict[str, Any]]:
    """Parse JSON payload supporting root arrays or collection objects."""
    try:
        text = data.decode("utf-8-sig")
        parsed = json.loads(text)
    except Exception as exc:
        raise SecurityValidationError(f"Malformed JSON: {exc}") from exc

    _check_json_depth(parsed)

    records: list[Any] = []
    if isinstance(parsed, list):
        records = parsed
    elif isinstance(parsed, dict):
        # Look for standard collection keys
        found_list = None
        for key in ("orders", "order_items", "products", "customers", "data", "items", "records", "rows", "results"):
            if isinstance(parsed.get(key), list):
                found_list = parsed[key]
                break
        if found_list is not None:
            records = found_list
        else:
            records = [parsed]
    else:
        raise SecurityValidationError("JSON root must be an array or an object")

    count = 0
    for item in records:
        if isinstance(item, dict):
            yield item
            count += 1
            if max_rows is not None and count >= max_rows:
                break


def parse_dataset_stream(
    file_type: str,
    data: bytes,
    *,
    max_rows: int | None = None,
) -> Iterator[dict[str, Any]]:
    """Unified entrypoint yielding row dictionaries from CSV, XLSX, or JSON."""
    ft = file_type.lower().strip(".")
    if ft == "csv":
        return parse_csv_stream(data, max_rows=max_rows)
    if ft == "xlsx":
        return parse_xlsx_stream(data, max_rows=max_rows)
    if ft == "json":
        return parse_json_stream(data, max_rows=max_rows)
    raise SecurityValidationError(f"Unsupported dataset format: {file_type}")
